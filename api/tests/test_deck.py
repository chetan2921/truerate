import importlib.util
import json
from pathlib import Path

import numpy as np
from pptx import Presentation
from test_audience import fake_embed, genuine_snapshot
from test_pipeline import FAKE_MODEL
from test_pricing import synthetic_rows

from truerate.pricing import validate
from truerate.signals import redteam

spec = importlib.util.spec_from_file_location("make_deck", Path(__file__).resolve().parents[1] / "scripts" / "make_deck.py")
make_deck = importlib.util.module_from_spec(spec)
spec.loader.exec_module(make_deck)


def test_deck_has_at_most_12_slides_with_the_real_numbers_and_no_handles(tmp_path):
    from truerate.experiment import run

    report = validate(synthetic_rows())
    rt = redteam([genuine_snapshot(i, np.random.default_rng(i)) for i in range(12)], FAKE_MODEL, fake_embed)
    rows = synthetic_rows(60, 7)
    fresh = [r | {"handle": f"fresh{i}", "holdout": False} for i, r in enumerate(synthetic_rows(12, 8))]
    experiments = json.loads(json.dumps(run([r for r in rows if not r["holdout"]], [r for r in rows if r["holdout"]], fresh,
                                            repeats=1, folds=3, points=("today", "ridge_v2"))))
    out = tmp_path / "deck.pptx"
    make_deck.build_deck(json.loads(json.dumps(report)), rt, None, out, experiments=experiments)
    deck = Presentation(out)
    assert 10 <= len(deck.slides) <= 12
    text = " ".join(shape.text_frame.text for s in deck.slides for shape in s.shapes if shape.has_text_frame)
    tables = " ".join(cell.text for s in deck.slides for shape in s.shapes if shape.has_table for row in shape.table.rows for cell in row.cells)
    assert f"{report['holdout']['model']['median_error']:.0%}" in tables  # the headline error comes from the report
    assert f"{report['holdout']['model']['within_2x']:.0%}" in tables  # and the scores judges know
    assert f"{rt['caught']['smart_fake']:.0%}" in tables
    assert f"{experiments['holdout']['ridge_v2/global']['error']:.0%}" in tables and "Ridge with the new features" in tables  # the models tried
    assert f"{experiments['coverage_by_width']['holdout']['2']:.0%}" in tables  # why the range is wide
    assert "Brand match" in tables
    assert "creator0" not in text + tables and "creator1" not in text + tables and "fresh0" not in text + tables
